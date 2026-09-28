#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import openpyxl
wb = openpyxl.load_workbook("/Users/kanyun/voice-ui-demos/bambi-lesson4-rowlevel.xlsx", data_only=True)
print("=== 工作表 ===", wb.sheetnames)
for ws in wb.worksheets:
    print(f"\n##### Sheet: {ws.title}  {ws.max_row} 行 × {ws.max_column} 列")
    for ri, row in enumerate(ws.iter_rows(min_row=1, max_row=min(6, ws.max_row), values_only=True), 1):
        print(f"  行{ri}:", [str(c)[:24] if c is not None else "" for c in row])
