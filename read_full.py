#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import openpyxl
wb = openpyxl.load_workbook("/Users/kanyun/voice-ui-demos/bambi-lesson4-data.xlsx", data_only=True)

print("############ 说明页全文 ############")
ws = wb["说明"]
for row in ws.iter_rows(values_only=True):
    cells = [str(c) for c in row if c is not None and str(c).strip()]
    if cells:
        print(" | ".join(cells))

print("\n############ 15题总览 全部 ############")
ws = wb["15题总览"]
for row in ws.iter_rows(values_only=True):
    print([str(c) if c is not None else "" for c in row])

print("\n############ 15字长答界 全部 ############")
ws = wb["15字长答界"]
for row in ws.iter_rows(values_only=True):
    cells = [str(c) for c in row if c is not None and str(c).strip()]
    if cells:
        print(" | ".join(cells))
