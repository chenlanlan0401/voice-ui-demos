#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""审视现有论证的逻辑漏洞：用原始聚合数据验证几个怀疑点。"""
import openpyxl
wb = openpyxl.load_workbook("/Users/kanyun/voice-ui-demos/bambi-lesson4-data.xlsx", data_only=True)

# ---- 读 15题总览 ----
ws = wb["15题总览"]
rows = list(ws.iter_rows(min_row=2, values_only=True))
overview = []
for r in rows:
    overview.append({
        "qid": r[0], "sub": r[1], "type": r[2], "title": str(r[3]),
        "n": r[4], "slow": float(r[5]), "short": float(r[6]),
        "long": float(r[7]), "thr": float(r[8]), "pass": r[9], "cls": r[10]
    })

print("========== 验证1：长答占比 是否 ≡ 100 - 短答占比（是否只是字数镜像）==========")
allmatch=True
for o in overview:
    diff = abs(o["long"] - (100-o["short"]))
    if diff>0.15:
        allmatch=False
        print(f"  Q{o['qid']}-{o['sub']} 长答{o['long']} vs 100-短答{100-o['short']:.1f}  差{diff:.1f}")
print("  结论：", "长答占比 = 100 - 短答占比，完全由字数决定，不是独立维度" if allmatch
      else "存在不一致，长答占比可能另有口径")

print("\n========== 验证2：门槛是相对分位（循环论证）？重算第70百分位 ==========")
slows = sorted([o["slow"] for o in overview])
import statistics
# 第70百分位（线性插值）
def pct(vals, p):
    vals=sorted(vals); idx=(len(vals)-1)*p; lo=int(idx); hi=min(lo+1,len(vals)-1)
    return vals[lo]+(vals[hi]-vals[lo])*(idx-lo)
p70 = pct(slows,0.70)
print(f"  15题慢答占比：{[round(x,1) for x in slows]}")
print(f"  重算第70百分位 = {p70:.1f}%  （表里门槛=24.3%）")
n_pass = sum(1 for o in overview if o["slow"]>=24.3)
print(f"  → 过门槛题数 = {n_pass}/15 = {n_pass/15*100:.0f}%  "
      f"（相对门槛必然让约30%的题过关，与题本身快慢无关）")

print("\n========== 验证3：慢答占比 与 长答占比 的关系（两轴是否真能区分）==========")
# 过门槛的题里，看长答占比怎么分三类
print("  过门槛的题（慢答≥24.3%）按长答占比排：")
passed=[o for o in overview if o["slow"]>=24.3]
for o in sorted(passed,key=lambda x:-x["long"]):
    print(f"    Q{o['qid']}-{o['sub']} {o['type']:<8} 慢答{o['slow']:>5}% 长答{o['long']:>5}% → {o['cls']}")

print("\n========== 验证4：开口时长分布——5秒阈值对各题型是否统一合理 ==========")
ws2 = wb["开口时长分布"]
hdr = [str(c) for c in next(ws2.iter_rows(min_row=1,max_row=1,values_only=True))]
# 找列
def col(name):
    for i,h in enumerate(hdr):
        if h and name in h: return i
    return -1
ci_cum5 = col("5s前累计")
ci_step = col("掉台阶")
ci_slow = col("慢答占比")
d2=[]
for r in ws2.iter_rows(min_row=2,values_only=True):
    d2.append(r)
print(f"  各题「5秒前累计开口%」分布（越高=大多数人5秒内已开口，该题本就快）：")
cum_vals=[]
for r in d2:
    cum = r[ci_cum5]
    cum_vals.append(float(cum))
    print(f"    Q{r[0]}-{r[1]} 5秒前累计{cum}%  掉台阶={r[ci_step]}")
print(f"  5秒前累计% 跨度：{min(cum_vals):.0f}% ~ {max(cum_vals):.0f}%  （差 {max(cum_vals)-min(cum_vals):.0f}个百分点）")
step_vals=[r[ci_step] for r in d2]
print(f"  「掉台阶?」取值分布：{set(step_vals)} → 若全是'是'则该指标无区分度")

print("\n========== 验证5：按题型看慢答占比（5秒阈值是否对题型不公平）==========")
from collections import defaultdict
by_type=defaultdict(list)
for o in overview:
    by_type[o["type"]].append(o["slow"])
for t,vals in sorted(by_type.items(),key=lambda x:-statistics.mean(x[1])):
    print(f"  {t:<10} 平均慢答{statistics.mean(vals):>5.1f}%  n={len(vals)}  明细{[round(v,1) for v in vals]}")
