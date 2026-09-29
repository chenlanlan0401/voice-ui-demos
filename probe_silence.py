#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""探查:行级数据里能不能反推'沉默'(有声音但没识别出文字)。
先摸清列结构,再统计空文字/零字数/异常记录的分布。"""
import openpyxl
from collections import defaultdict

wb=openpyxl.load_workbook("/Users/kanyun/voice-ui-demos/bambi-lesson4-rowlevel.xlsx",data_only=True)
ws=wb["作答明细"]

# 1) 列头
hdr=[c.value for c in ws[1]]
print("列头:", hdr)
print("总行数(含表头):", ws.max_row)

# 2) 逐行看 秒数/字数/文字 的空值情况
rows=[]
for r in ws.iter_rows(min_row=2,values_only=True):
    rows.append(r)
print("数据行:", len(rows))

empty_txt=0; zero_w=0; none_w=0; both_empty=0
neg_or_zero_sec=0
samples_empty=[]
for r in rows:
    sec=r[4]; w=r[5]; txt=r[6]
    t=(str(txt).strip() if txt is not None else "")
    is_empty_txt = (t=="" or t.lower()=="none" or t=="nan")
    if is_empty_txt:
        empty_txt+=1
        if len(samples_empty)<20: samples_empty.append((r[0],r[1],sec,w,repr(txt)))
    try:
        wi=int(w) if w is not None else None
    except: wi=None
    if wi==0: zero_w+=1
    if w is None: none_w+=1
    if is_empty_txt and (wi==0 or wi is None): both_empty+=1
    try:
        s=float(sec) if sec is not None else None
        if s is not None and s<=0: neg_or_zero_sec+=1
    except: pass

print(f"\n空文字行: {empty_txt}  ({empty_txt/len(rows)*100:.1f}%)")
print(f"字数=0 行: {zero_w}")
print(f"字数=None 行: {none_w}")
print(f"文字空 且 字数0/None: {both_empty}")
print(f"秒数<=0 行: {neg_or_zero_sec}")

print("\n=== 空文字样本(题号1/题号2/秒/字/原文) ===")
for s in samples_empty:
    print("  ",s)

# 3) 每题的作答人数是否一致?若某些学生某题无记录,说明'没开口'可能直接缺行
cnt=defaultdict(int)
for r in rows:
    cnt[f"{r[0]}-{r[1]}"]+=1
print("\n=== 每题记录数(看是否等人数,缺行=可能没开口没进表) ===")
for k in sorted(cnt): print(f"  {k}: {cnt[k]}")
ns=list(cnt.values())
print(f"  min={min(ns)} max={max(ns)} 差={max(ns)-min(ns)}")
