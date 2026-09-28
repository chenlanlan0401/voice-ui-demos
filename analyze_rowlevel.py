#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
题级主表 + 权衡曲线（方案中立）。
口径：
- 慢答者 = 开口秒数 > SLOW_SEC（默认埋点原始 5s，可复现报告 54%）
- 长答 = 回答字数 >= LONG_CHARS（15字，沿用已验证的长答界）
- 慢答者长答率 = 慢答者里 长答 的比例（条件指标，真第二轴）
- 误伤率 = (慢答 AND 长答) 学生 / 全部作答学生
- 秒数三档 = 慢答者开口秒数的 中位/P75/P90
- 高精准判定：慢答者占比≥SLOW_SHARE_THR 且 慢答者长答率≥COND_LONG_THR → 该等
             慢答者占比≥门槛 但 慢答者长答率低 → 答案本就短
             慢答者占比<门槛 → 立即开麦
"""
import openpyxl, statistics, json
from collections import defaultdict

SLOW_SEC = 5.0          # 慢答阈值（埋点原始，含1-2s延迟）
LONG_CHARS = 15         # 长答界
SLOW_SHARE_THR = 0.25   # 慢答者占比门槛（绝对值，不用相对分位）
COND_LONG_THR = 0.55    # 慢答者长答率门槛（高精准：慢的人里过半在憋长答才算该等）

wb = openpyxl.load_workbook("/Users/kanyun/voice-ui-demos/bambi-lesson4-rowlevel.xlsx", data_only=True)
ws = wb["作答明细"]
# 题型映射（从旧总览表拿）
wb2 = openpyxl.load_workbook("/Users/kanyun/voice-ui-demos/bambi-lesson4-data.xlsx", data_only=True)
qtype={}
for r in wb2["15题总览"].iter_rows(min_row=2, values_only=True):
    qtype[(str(r[0]),str(r[1]))]=r[2]

rows=defaultdict(list)  # (qid,sub) -> list of (sec, words, title)
for r in ws.iter_rows(min_row=2, values_only=True):
    qid,sub,title,sid,sec,words,txt = r
    rows[(str(qid),str(sub))].append((float(sec),int(words),str(title)))

def pct(vals,p):
    vals=sorted(vals);
    if not vals: return 0
    idx=(len(vals)-1)*p; lo=int(idx); hi=min(lo+1,len(vals)-1)
    return vals[lo]+(vals[hi]-vals[lo])*(idx-lo)

table=[]
for key,recs in rows.items():
    n=len(recs)
    title=recs[0][2]
    slow=[(s,w) for s,w in [(x[0],x[1]) for x in recs] if s>SLOW_SEC]
    n_slow=len(slow)
    slow_share=n_slow/n
    slow_long=[ (s,w) for s,w in slow if w>=LONG_CHARS ]
    cond_long_rate = len(slow_long)/n_slow if n_slow>0 else 0
    misfire = len(slow_long)/n   # 误伤率：慢AND长 / 全部
    slow_secs=[s for s,w in slow]
    med = pct(slow_secs,0.5) if slow_secs else 0
    p75 = pct(slow_secs,0.75) if slow_secs else 0
    p90 = pct(slow_secs,0.90) if slow_secs else 0
    # 旧尺：两步走（先卡慢答占比，再看慢答者长答率）
    if slow_share>=SLOW_SHARE_THR and cond_long_rate>=COND_LONG_THR:
        verdict_old="该等"
    elif slow_share>=SLOW_SHARE_THR:
        verdict_old="答案本就短"
    else:
        verdict_old="立即开麦"
    table.append({
        "qid":key[0],"sub":key[1],"title":title[:20],
        "type":qtype.get(key,"?"),"n":n,
        "slow_share":round(slow_share*100,1),
        "cond_long_rate":round(cond_long_rate*100,1),
        "misfire":round(misfire*100,1),
        "med":round(med,1),"p75":round(p75,1),"p90":round(p90,1),
        "verdict_old":verdict_old,
    })

table.sort(key=lambda x:-x["misfire"])

# --- 新尺：找误伤率的自然断层，作为绝对门槛 ---
mf=[t["misfire"] for t in table]
gaps=[(round(mf[i-1]-mf[i],1), (mf[i-1]+mf[i])/2, i) for i in range(1,len(mf))]
biggest=max(gaps, key=lambda x:x[0])
MISFIRE_THR = round(biggest[1],1)   # 断层中点做门槛
print(f"误伤率降序断层：最大跳变 {biggest[0]} 个百分点，落在 {mf[biggest[2]-1]}% → {mf[biggest[2]]}% 之间")
print(f"→ 新尺门槛(误伤率) = {MISFIRE_THR}%（数据自然断层，非拍脑袋）\n")

for t in table:
    t["verdict_new"]="该等" if t["misfire"]>=MISFIRE_THR else "立即开麦"
    t["diff"] = "★分歧" if (t["verdict_new"]=="该等")!=(t["verdict_old"]=="该等") else ""

print(f"参数: 慢答>{SLOW_SEC}s | 长答>={LONG_CHARS}字 | 旧尺门槛[慢答{SLOW_SHARE_THR*100:.0f}%+慢答者长答率{COND_LONG_THR*100:.0f}%] | 新尺门槛[误伤率{MISFIRE_THR}%]")
print("="*128)
print(f"{'题号-问':<11}{'题型':<9}{'人数':>4} {'慢答%':>6} {'慢长率%':>8} {'误伤%':>6} {'中位s':>6}{'P75':>5}{'P90':>5}  {'旧尺':<8}{'新尺':<8}{'分歧':<6} 题目")
print("-"*128)
for t in table:
    print(f"{t['qid']}-{t['sub']:<5}{t['type']:<9}{t['n']:>4} {t['slow_share']:>6} {t['cond_long_rate']:>8} "
          f"{t['misfire']:>6} {t['med']:>6}{t['p75']:>5}{t['p90']:>5}  {t['verdict_old']:<8}{t['verdict_new']:<8}{t['diff']:<6} {t['title']}")

no=sum(1 for t in table if t['verdict_old']=='该等')
nn=sum(1 for t in table if t['verdict_new']=='该等')
nd=sum(1 for t in table if t['diff'])
print("-"*128)
print(f"旧尺判该等 {no} 题 | 新尺判该等 {nn} 题 | 分歧 {nd} 题")

with open("/Users/kanyun/voice-ui-demos/rowlevel_result.json","w",encoding="utf-8") as f:
    json.dump({"params":{"slow_sec":SLOW_SEC,"long_chars":LONG_CHARS,
               "slow_share_thr":SLOW_SHARE_THR,"cond_long_thr":COND_LONG_THR,
               "misfire_thr":MISFIRE_THR},
               "table":table},f,ensure_ascii=False)
