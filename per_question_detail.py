#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""每道题的完整画像:题目全文 + 指标 + 慢答者真实原话样本,供人工逐题判断。"""
import openpyxl, json
from collections import defaultdict

SLOW_SEC=5.0; LONG_CHARS=15
wb = openpyxl.load_workbook("/Users/kanyun/voice-ui-demos/bambi-lesson4-rowlevel.xlsx", data_only=True)
ws = wb["作答明细"]
wb2 = openpyxl.load_workbook("/Users/kanyun/voice-ui-demos/bambi-lesson4-data.xlsx", data_only=True)
qtype={}
for r in wb2["15题总览"].iter_rows(min_row=2, values_only=True):
    qtype[(str(r[0]),str(r[1]))]=r[2]

recs=defaultdict(list)
for r in ws.iter_rows(min_row=2, values_only=True):
    qid,sub,title,sid,sec,words,txt=r
    recs[(str(qid),str(sub))].append({"sec":float(sec),"w":int(words),"txt":str(txt or "")})

def pct(vals,p):
    vals=sorted(vals)
    if not vals: return 0
    idx=(len(vals)-1)*p; lo=int(idx); hi=min(lo+1,len(vals)-1)
    return vals[lo]+(vals[hi]-vals[lo])*(idx-lo)

out=[]
for key,rows in recs.items():
    n=len(rows)
    title=next((str(r[3]) for r in wb2["15题总览"].iter_rows(min_row=2,values_only=True)
                if str(r[0])==key[0] and str(r[1])==key[1]), "")
    slow=[r for r in rows if r["sec"]>SLOW_SEC]
    slow_long=[r for r in slow if r["w"]>=LONG_CHARS]
    slow_short=[r for r in slow if r["w"]<LONG_CHARS]
    ss=[r["sec"] for r in slow]
    # 慢答长答样本:按字数从多到少取5条(看"憋出来的长句"长啥样)
    samp_long=sorted(slow_long,key=lambda x:-x["w"])[:5]
    # 慢答短答样本:取5条(看"慢但没说啥"长啥样)
    samp_short=sorted(slow_short,key=lambda x:x["w"])[:5]
    out.append({
        "qid":key[0],"sub":key[1],"type":qtype.get(key,"?"),"title":title,"n":n,
        "slow_share":round(len(slow)/n*100,1),
        "cond_long":round(len(slow_long)/len(slow)*100,1) if slow else 0,
        "misfire":round(len(slow_long)/n*100,1),
        "med":round(pct(ss,0.5),1) if ss else 0,
        "p75":round(pct(ss,0.75),1) if ss else 0,
        "p90":round(pct(ss,0.90),1) if ss else 0,
        "n_slow":len(slow),"n_slowlong":len(slow_long),
        "samp_long":[{"sec":s["sec"],"w":s["w"],"txt":s["txt"][:60]} for s in samp_long],
        "samp_short":[{"sec":s["sec"],"w":s["w"],"txt":s["txt"][:60]} for s in samp_short],
    })

out.sort(key=lambda x:-x["misfire"])
with open("/Users/kanyun/voice-ui-demos/per_question.json","w",encoding="utf-8") as f:
    json.dump(out,f,ensure_ascii=False)

# 控制台速览
for q in out:
    print(f"\n{'='*70}")
    print(f"Q{q['qid']}-{q['sub']} [{q['type']}] {q['title']}")
    print(f"  作答{q['n']}人 | 慢答{q['slow_share']}% | 慢答者长答率{q['cond_long']}% | 误伤{q['misfire']}% | 慢答者开口 中位{q['med']}s/P75={q['p75']}s")
    print(f"  --- 慢答者【长答】样本(字数降序) ---")
    for s in q['samp_long'][:3]:
        print(f"    {s['sec']}s/{s['w']}字: {s['txt']}")
    print(f"  --- 慢答者【短答】样本 ---")
    for s in q['samp_short'][:3]:
        print(f"    {s['sec']}s/{s['w']}字: {s['txt']}")
print(f"\n已生成 per_question.json")
