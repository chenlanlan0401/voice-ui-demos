#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import openpyxl, statistics
from collections import defaultdict, Counter
wb = openpyxl.load_workbook("/Users/kanyun/voice-ui-demos/bambi-lesson4-rowlevel.xlsx", data_only=True)

print("########## 说明页全文 ##########")
for row in wb["说明"].iter_rows(values_only=True):
    cells=[str(c) for c in row if c is not None and str(c).strip()]
    if cells: print(" | ".join(cells))

ws = wb["作答明细"]
rows=[]
bad_sec=0; bad_word=0
for r in ws.iter_rows(min_row=2, values_only=True):
    qid,sub,title,sid,sec,words,txt = r
    if sec is None: bad_sec+=1
    if words is None: bad_word+=1
    rows.append({"qid":qid,"sub":sub,"title":str(title),"sid":sid,
                 "sec":sec,"words":words,"txt":txt})

print(f"\n########## 数据质量体检 ##########")
print(f"总记录 {len(rows)}  | 开口秒数缺失 {bad_sec} | 字数缺失 {bad_word}")
secs=[r["sec"] for r in rows if isinstance(r["sec"],(int,float))]
words=[r["words"] for r in rows if isinstance(r["words"],(int,float))]
print(f"开口秒数 范围 {min(secs):.1f}~{max(secs):.1f}  中位 {statistics.median(secs):.1f}  均值 {statistics.mean(secs):.1f}")
print(f"回答字数 范围 {min(words)}~{max(words)}  中位 {statistics.median(words)}  均值 {statistics.mean(words):.1f}")
# 异常值检查
neg=[s for s in secs if s<0]; huge=[s for s in secs if s>60]
print(f"负开口秒数 {len(neg)} 条 | 开口>60秒 {len(huge)} 条")
zero_word=sum(1 for w in words if w==0)
print(f"字数=0(没说话?) {zero_word} 条 = {zero_word/len(words)*100:.1f}%")

# 每题作答人数
byq=defaultdict(int)
for r in rows: byq[(r["qid"],r["sub"])]+=1
print(f"\n共 {len(byq)} 道题(题号+小问)  每题作答人数:")
for k,v in sorted(byq.items()): print(f"  {k[0]}-{k[1]}: {v}人")

# 学生数
sids=set(r["sid"] for r in rows)
print(f"\n去重学生数(chatid): {len(sids)}")
