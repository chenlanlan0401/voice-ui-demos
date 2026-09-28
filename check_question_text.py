#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""核实：数据里有没有 AI 老师问的完整问题原文,还是只有简化题干。"""
import openpyxl
# 行级表
wb = openpyxl.load_workbook("/Users/kanyun/voice-ui-demos/bambi-lesson4-rowlevel.xlsx", data_only=True)
print("### 行级表 sheet:", wb.sheetnames)
ws=wb["作答明细"]
hdr=[str(c) for c in next(ws.iter_rows(min_row=1,max_row=1,values_only=True))]
print("### 行级表列头:", hdr)
# 看 title 列有几种不同取值,是否就是简化题干
titles={}
for r in ws.iter_rows(min_row=2,values_only=True):
    key=(str(r[0]),str(r[1]))
    titles.setdefault(key,set()).add(str(r[2]))
print(f"### 共 {len(titles)} 道题,每题 title 取值:")
for k,v in sorted(titles.items()):
    for t in v:
        print(f"  {k[0]}-{k[1]}: 【{t}】(长度{len(t)})")

# 说明页看有没有提到问题原文来源
print("\n### 说明页:")
for row in wb["说明"].iter_rows(values_only=True):
    cells=[str(c) for c in row if c is not None and str(c).strip()]
    if cells: print("  "+" | ".join(cells))

# 汇总表也看一眼 title
print("\n### 汇总表 15题总览 的题目列:")
wb2=openpyxl.load_workbook("/Users/kanyun/voice-ui-demos/bambi-lesson4-data.xlsx",data_only=True)
for r in wb2["15题总览"].iter_rows(min_row=2,values_only=True):
    print(f"  {r[0]}-{r[1]}: 【{r[3]}】(长度{len(str(r[3]))})")
