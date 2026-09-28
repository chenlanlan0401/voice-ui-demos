#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""lesson4 最终判定:误伤率分三档,模糊带调出语义+原话供自主判定。"""
import openpyxl, json
from collections import defaultdict

SLOW_SEC=5.0; LONG_CHARS=15
# 三档门槛(基于数据自然断层):>=25 明显该等; <=10 明显不用; 中间为模糊带
HI=25.0; LO=10.0

wb = openpyxl.load_workbook("/Users/kanyun/voice-ui-demos/bambi-lesson4-rowlevel.xlsx", data_only=True)
ws = wb["作答明细"]
recs=defaultdict(list)
for r in ws.iter_rows(min_row=2, values_only=True):
    qid,sub,title,sid,sec,words,txt=r
    recs[(str(qid),str(sub))].append({"sec":float(sec),"w":int(words),"txt":str(txt or "")})

def pct(vals,p):
    vals=sorted(vals)
    if not vals: return 0
    idx=(len(vals)-1)*p; lo=int(idx); hi=min(lo+1,len(vals)-1)
    return vals[lo]+(vals[hi]-vals[lo])*(idx-lo)

rows=[]
for key,rs in recs.items():
    n=len(rs)
    slow=[x for x in rs if x["sec"]>SLOW_SEC]
    slow_long=[x for x in slow if x["w"]>=LONG_CHARS]
    slow_short=[x for x in slow if x["w"]<LONG_CHARS]
    ss=[x["sec"] for x in slow]
    mf=round(len(slow_long)/n*100,1)
    cl=round(len(slow_long)/len(slow)*100,1) if slow else 0
    band="明显该等" if mf>=HI else ("明显不用" if mf<=LO else "模糊带")
    rows.append({"key":f"{key[0]}-{key[1]}","n":n,
        "slow_share":round(len(slow)/n*100,1),"cond_long":cl,"misfire":mf,
        "n_slow":len(slow),"med":round(pct(ss,0.5),1) if ss else 0,
        "p75":round(pct(ss,0.75),1) if ss else 0,"band":band,
        "samp_long":sorted(slow_long,key=lambda x:-x["w"])[:4],
        "samp_short":sorted(slow_short,key=lambda x:x["w"])[:4],
    })

rows.sort(key=lambda x:-x["misfire"])
# 题目文本映射
qtitle={}
for r in ws.iter_rows(min_row=2,values_only=True):
    qtitle[f"{r[0]}-{r[1]}"]=str(r[2])

print("="*70)
for r in rows:
    print(f"\n[{r['band']}] {r['key']} 误伤{r['misfire']}% | 慢答{r['slow_share']}%({r['n_slow']}人) | 慢答者长答率{r['cond_long']}% | P75={r['p75']}s")
    print(f"  题目: {qtitle[r['key']]}")
    if r['band']=="模糊带":
        print(f"  【长答样本】")
        for s in r['samp_long'][:3]:
            print(f"    {s['sec']}s/{s['w']}字: {s['txt'][:50]}")
        print(f"  【短答样本】")
        for s in r['samp_short'][:3]:
            print(f"    {s['sec']}s/{s['w']}字: {s['txt'][:50]}")

# 存json
for r in rows:
    r['title']=qtitle[r['key']]
    r['samp_long']=[{"sec":s["sec"],"w":s["w"],"txt":s["txt"][:80]} for s in r['samp_long']]
    r['samp_short']=[{"sec":s["sec"],"w":s["w"],"txt":s["txt"][:80]} for s in r['samp_short']]
json.dump(rows,open("/Users/kanyun/voice-ui-demos/final_lesson4.json","w"),ensure_ascii=False)
print("\n\n已存 final_lesson4.json")
n_band=defaultdict(int)
for r in rows: n_band[r['band']]+=1
print("分档:",dict(n_band))
